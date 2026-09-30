# SPDX-FileCopyrightText: Copyright (c) 2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""What the file tools do to the workspace, run where the files are.

In the sandbox, whose mounts are the boundary: the host never opens a path the agent chose, so a
link planted from the shell between a check and an open reaches nothing there. The standard library
only at the top, since this source is sent to the container and run as it is; without any sandbox
the tools import it here instead. Each operation returns what its tool says to the model.
"""
import base64
import contextlib
import csv
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

# Where the sandbox mounts the workspace, as upstream's start_local_sandbox.sh does.
SANDBOX_ROOT = "/workspace"
CENSUS_ROWS = 40
_NOISE = {".git", "node_modules", ".venv", ".mypy_cache", ".pytest_cache"}
_PART_OF = ("\n\n[only the first {kept} {unit} were read, of {whole}; the rest is not shown and searching this "
            "file will not find it]")
_OFFICE = {".docx", ".xlsx", ".pptx", ".doc", ".xls", ".ppt"}
_DOCUMENTS = _OFFICE | {".pdf"}
# On disk, not in memory: every operation is a fresh process, and a second search over the same tree
# re-parsed 1950 PDFs from scratch when nothing was kept.
_CACHE = Path(tempfile.gettempdir()) / f"nat-extract-{os.getuid()}"

# Out of process: one malformed PDF can hang pdfminer or crash the interpreter, uncatchably.
_PDF_CHILD = ("import sys, pdfplumber\n"
              "with pdfplumber.open(sys.argv[1]) as pdf:\n"
              "    sys.stdout.write('\\n'.join((p.extract_text() or '') for p in pdf.pages[:40]))\n"
              # Said, not guessed: without the count the parent cannot tell 40 pages from 400,
              # and read_file then reports the head of a long document as the whole of it.
              "    sys.stderr.write(str(len(pdf.pages)))\n")


def _allowed(root: Path, p: Path, data: str, write: bool = False) -> bool:
    # A staged world may read through its links into the data its sandbox mounts, never write there.
    real = p.resolve()
    return real.is_relative_to(root) or (not write and bool(data) and real.is_relative_to(Path(data).resolve()))


def resolve(root: str, rel: str, data: str = "", write: bool = False) -> Path:
    # Briefs echo the root path whole, partial or not at all; each form folds under the root.
    root = Path(root).resolve()
    given = Path(rel.strip())
    parts = tuple(x for x in given.parts if x not in ("/", "", "."))
    if given.is_absolute() or rel.strip().startswith("~"):
        if Path(os.path.realpath(given)).is_relative_to(root):
            parts = Path(os.path.realpath(given)).relative_to(root).parts
        elif given.is_relative_to(SANDBOX_ROOT):
            parts = given.relative_to(SANDBOX_ROOT).parts
        else:
            raise ValueError(f"{rel} is outside the workspace, and nothing outside it is kept after the run; "
                             f"name it relative to the workspace root")
    else:
        parts = parts[next((k for k in range(len(root.parts) - 1, 0, -1) if parts[:k] == root.parts[-k:]), 0):]
    p = Path(os.path.normpath(root.joinpath(*parts)))
    if not p.is_relative_to(root) or not _allowed(root, p, data) or (
            write and p != root and not _allowed(root, p.parent, data, write=True)):
        raise ValueError(f"path escapes workspace: {rel}")
    return p


def nothing_read(path: str, size: int, offset: int) -> str:
    """An empty read means an empty file or an offset past the end, and the bare "" it returned
    means those and a broken tool alike -- one run asked for the same empty file twice over."""
    if size <= 0:
        return f"{path} exists and is empty; there is nothing in it to read."
    return (f"nothing at offset {offset}: {path} is {size} characters long, so that is past its "
            f"end. Read it from 0, or from an offset below {size}.")


def _pdf_text(p: Path) -> str | None:
    # A workspace tree carries empty placeholder files; an empty one is not a broken PDF.
    if p.stat().st_size == 0:
        return None
    try:
        done = subprocess.run([sys.executable, "-c", _PDF_CHILD, str(p)],
                              capture_output=True, timeout=60, check=False)
    except subprocess.TimeoutExpired:
        return None
    if done.returncode != 0:
        return None
    text = done.stdout.decode("utf-8", "ignore").strip() or None
    # The child wrote the page census to stderr: 40 pages of a 266-page filing is not the filing,
    # and read_file reports a head it was never told was a head as the whole document.
    if text:
        with contextlib.suppress(ValueError):
            whole = int(done.stderr.decode("utf-8", "ignore").strip() or 0)
            if whole > 40:
                text += _PART_OF.format(kept=40, whole=whole, unit="pages of this PDF")
    return text


def extract(p: Path) -> str | None:
    """Text from a workspace file, or None: Office formats are zipped XML read via the stdlib;
    reading them as UTF-8 yields mojibake that floods the context window. A file that could not be
    parsed is kept as None too: each attempt at a broken PDF costs the full timeout."""
    try:
        st = p.stat()
    except OSError:
        return None
    kept = _CACHE / hashlib.sha1(f"{p.resolve()}|{st.st_size}|{st.st_mtime_ns}".encode()).hexdigest()
    with contextlib.suppress(OSError, ValueError):
        return json.loads(kept.read_text(encoding="utf-8"))
    text = _extract_uncached(p)
    with contextlib.suppress(OSError):
        _CACHE.mkdir(parents=True, exist_ok=True)
        kept.write_text(json.dumps(text), encoding="utf-8")
    return text


def _extract_uncached(p: Path) -> str | None:
    import zipfile

    suffix = p.suffix.lower()
    try:
        with p.open("rb") as fh:
            head = fh.read(8)
    except OSError:
        return None
    # By content, not name: these worlds hold .docx that are CSV, .doc that are docx, a pdf named .docx.
    # The zip's own header, not is_zipfile: an old binary that embeds a zip has its directory near the end.
    office = suffix in _OFFICE
    if office and head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(p) as z:
                skip = ("docProps/", "theme", "styles", "settings", "fontTable", "Content_Types",
                        "slideLayout", "slideMaster", "notesMaster")
                parts = [n for n in z.namelist() if n.endswith(".xml") and "rels" not in n
                         and not any(s in n for s in skip)]
                chunks = []
                for name in parts[:40]:
                    raw = z.read(name).decode("utf-8", errors="ignore")
                    chunks.append(re.sub(r"<[^>]+>", " ", raw))
                text = re.sub(r"\s+", " ", " ".join(chunks)).strip() or None
                # A document read down to its first 40 parts is not the document, and silence here
                # reads downstream as "this is all of it".
                if text and len(parts) > 40:
                    text += _PART_OF.format(kept=40, whole=len(parts), unit="parts of this file")
                return text
        except Exception:
            return None
    if office and head.startswith(b"\xd0\xcf\x11\xe0"):
        # The pre-2007 binaries, through catdoc's readers: in the image, or beside this python.
        beside = os.pathsep.join([os.path.dirname(sys.executable), os.environ.get("PATH", "")])
        tool = shutil.which({"doc": "catdoc", "ppt": "catppt", "xls": "xls2csv"}[suffix[1:4]], path=beside)
        if not tool:
            return None
        try:
            done = subprocess.run([tool, "-d", "utf-8", str(p)], capture_output=True, encoding="utf-8",
                                  errors="replace", timeout=60, check=False)
        except (OSError, subprocess.TimeoutExpired):
            return None
        return done.stdout.strip() or None
    if suffix == ".pdf" or head.startswith(b"%PDF"):
        text = _pdf_text(p)
        # An unparseable .pdf that is really plain text is an agent-written deliverable: let it read back.
        if text is not None:
            return text
    # ASCII-leading binaries still read as mojibake, so the extension decides, not a byte probe.
    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".zip", ".bin", ".so"}:
        return None
    try:
        data = p.read_bytes()
    except Exception:
        return None
    if b"\x00" in data[:4096]:
        return None
    # GB18030 next, strictly: a Chinese office file is often GBK, and detection misreads short ones as Korean.
    for encoding in ("utf-8", "gb18030"):
        with contextlib.suppress(UnicodeDecodeError):
            return data.decode(encoding)
    return data.decode("utf-8", errors="replace")


def formatted(root: str, pairs, contains: str, max_entries: int) -> str:
    """One listing format for both trees, the sandbox's and a task container's: the files, or past the
    cap a folder census, which is the only thing an arbitrary slice of them could not tell."""
    needle = (contains or "").strip().lower()
    hits: list[str] = []
    census: dict[str, int] = {}
    for rel, size in pairs:
        if not rel or (needle and needle not in rel.lower()):
            continue
        census[str(Path(rel).parent)] = census.get(str(Path(rel).parent), 0) + 1
        hits.append(f"{rel}  ({size} bytes)" if size else rel)
    # The absolute root, so code run in a sandbox can open these files by path.
    head = f"workspace root: {root}"
    if not hits:
        return head + ("\n(no file matches %r)" % contains if needle else "\n(empty)")
    if len(hits) <= max_entries:
        return "\n".join([head, *hits])
    rows = sorted(census.items(), key=lambda kv: -kv[1])[:CENSUS_ROWS]
    return (f"{head}\n{len(hits)} files match -- too many to list. Folders, largest first; "
            f"open one with `subdir`, or filter with `contains`:\n" +
            "\n".join(f"{d}/  ({n} files)" for d, n in rows))


def _shown(root: Path, p: Path, staged: bool, data: str) -> bool:
    """A staged world is walked whole; one the user named drops its metadata, .git alone filling
    a listing, and whatever its links reach outside."""
    return (staged or not _NOISE & set(p.relative_to(root).parts)) and _allowed(root, p, data)


def listing(root: str, subdir: str = "", contains: str = "", max_entries: int = 200, staged: bool = False,
            data: str = "") -> str:
    top = Path(root).resolve()
    base = resolve(root, subdir, data) if subdir else top
    if not base.is_dir():
        return f"not a directory: {subdir}"
    pairs = [(str(p.relative_to(top)), p.stat().st_size)
             for p in sorted(base.rglob("*")) if p.is_file() and _shown(top, p, staged, data)]
    # The path code in the sandbox opens, not the host's.
    return formatted(SANDBOX_ROOT, pairs, contains, max_entries)


def read(root: str, path: str, offset: int = 0, max_chars: int = 20000, data: str = "") -> str:
    p = resolve(root, path, data)
    if p.is_dir():
        return f"{path} is a directory -- list it with list_directory, or name a file in it."
    if not p.is_file():
        return f"no such file: {path}"
    text = extract(p)
    if text is None:
        image = p.suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"}
        seen = " Look at it with view_image." if image else ""
        return f"{path} ({p.stat().st_size} bytes) has no text this tool can extract.{seen}"
    chunk = text[offset:offset + max_chars]
    if not chunk:
        return nothing_read(path, len(text), offset)
    rest = len(text) - offset - len(chunk)
    if rest <= 0:
        return chunk
    # Say where to continue, or the model re-reads the same head and looks like a repeat loop.
    return f"{chunk}\n... {rest} more characters, call again with offset={offset + len(chunk)}"


def image(root: str, path: str, data: str = "") -> dict:
    """The image as the model is sent it, a PNG at a size it reads well, with the line that says so."""
    p = resolve(root, path, data)
    if not p.is_file():
        return {"text": f"no such file: {path}"}
    try:
        from PIL import Image
        with Image.open(p) as im:
            im = im.convert("RGB")
            w, h = im.size
            # Below ~1280px on the long side small text is misread; past 1600 only the token count grows.
            k = min(max(1280, max(w, h)), 1600) / max(w, h)
            if k != 1:
                im = im.resize((max(1, round(w * k)), max(1, round(h * k))))
            buf = io.BytesIO()
            im.save(buf, "PNG")
    except Exception:
        return {"text": f"{path} is not an image this tool can open (png, jpg, gif, bmp, webp)."}
    return {"text": f"{p.relative_to(Path(root).resolve())} ({w}x{h})",
            "png": base64.b64encode(buf.getvalue()).decode()}


def _add_table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows]
    # The `|---|:--:|` rule under a markdown header carries no data, so it must not become a row.
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    if not cells:
        return
    table = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            table.cell(i, j).text = value


def _write_docx(p: Path, text: str) -> str:
    """Build a real Word document out of markdown-ish text."""
    from docx import Document

    doc = Document()
    para: list[str] = []
    rows: list[str] = []

    def flush_para() -> None:
        if para:
            doc.add_paragraph(" ".join(para))
            para.clear()

    def flush_rows() -> None:
        if rows:
            _add_table(doc, rows)
            rows.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith("|") and line.endswith("|"):
            flush_para()
            rows.append(line)
            continue
        flush_rows()
        head = re.match(r"(#{1,3})\s+(.+)", line)
        bullet = re.match(r"([-*+•]|\d+[.)])\s+", line)
        if head or bullet or not line:
            flush_para()
        if head:
            doc.add_heading(head[2].strip(" #"), level=len(head[1]))
        elif bullet:
            doc.add_paragraph(line)
        elif line:
            para.append(line)
    flush_rows()
    flush_para()
    doc.save(p)
    return f"{len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables"


def _write_xlsx(p: Path, text: str) -> str:
    """Build a real Excel workbook out of CSV text."""
    from openpyxl import Workbook

    def _typed(cell: str):
        """A number written as text is one Excel flags and every formula skips, and a column holding
        both sorts worse than one holding either. Plain decimals convert; a leading zero, a plus or
        an exponent means an identifier -- 007, +1, 1e5 -- and stays the string it was sent as."""
        body = cell.strip()
        if not re.fullmatch(r"-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?", body):
            return cell
        return float(body) if "." in body else int(body)

    wb = Workbook()
    ws = wb.active
    # StringIO rather than splitlines(): only a real stream keeps a newline inside a quoted field.
    for row in csv.reader(io.StringIO(text)):
        ws.append([_typed(c) for c in row])
    wb.save(p)
    return f"{ws.max_row} rows x {ws.max_column} columns"


def _write_pptx(p: Path, text: str) -> str:
    """Build a real PowerPoint deck, one slide per blank-line-separated block."""
    from pptx import Presentation

    prs = Presentation()
    layout = prs.slide_layouts[1]
    for block in re.split(r"\n[ \t]*\n", text.strip()):
        lines = [ln.strip().lstrip("#-*• ").strip() for ln in block.splitlines() if ln.strip()]
        if not lines:
            continue
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = lines[0]
        slide.placeholders[1].text = "\n".join(lines[1:])
    prs.save(p)
    return f"{len(prs.slides)} slides"


BUILDERS = {".docx": _write_docx, ".xlsx": _write_xlsx, ".pptx": _write_pptx}


def built(target: Path, content: str) -> str:
    """Build `content` into `target` in the shape its extension asks for. -> what was built."""
    build = BUILDERS.get(target.suffix.lower())
    if build is None:
        target.write_text(content, encoding="utf-8")
        return f"{content.count(chr(10)) + 1} lines"
    return build(target, content)


def write(root: str, path: str, content: str, data: str = "") -> str:
    p = resolve(root, path, data, write=True)
    # An empty or directory `path` writes onto the directory itself.
    if p == Path(root).resolve() or p.is_dir():
        return f"`path` must name a file inside the workspace; {path!r} names a directory."
    p.parent.mkdir(parents=True, exist_ok=True)
    # Writing to a symlink never means editing its target; linked worlds would reject or corrupt.
    if p.is_symlink():
        p.unlink()
    shown = p.relative_to(Path(root).resolve())
    build = BUILDERS.get(p.suffix.lower())
    if build is None:
        # Said at the moment the lines go: a description telling the agent to prefer
        # edit_file moved 1 framework in 4, and the loss is silent otherwise.
        before = p.read_text(encoding="utf-8", errors="ignore").count("\n") + 1 if p.is_file() else 0
        try:
            p.write_text(content, encoding="utf-8")
        except OSError as exc:
            return f"could not write {path}: {exc.strerror or exc}."
        after = content.count("\n") + 1
        note = (f" It held {before} lines and now holds {after}; the other {before - after} are "
                "gone. If that was not intended, edit_file changes one passage and leaves "
                "the rest." if before > after + max(5, before // 10) else "")
        return f"wrote {shown} ({len(content)} chars){note}"
    try:
        detail = build(p, content)
    except Exception as exc:
        # Text saved under an Office name looks delivered and grades zero; the failure must be heard.
        return f"failed to build {p.name}: {exc}. Resend `content` in the shape that extension expects."
    return f"wrote {shown} ({detail})"


def literal(query: str) -> str:
    """Why a regex found nothing: the search is plain text. Said in the description and ignored --
    7 of nemotron's 18 searches were regexes, 14 empty -- so it goes in the answer."""
    return (" -- this search is plain text, so those regex characters had to appear literally; "
            "for a pattern use bash `grep -rnE`." if re.search(r"[.*+?\[\]\\^$|()]", query) else "")


def grep(root: str, query: str, subdir: str = "", path_contains: str = "", max_hits: int = 60,
         max_seconds: float = 90.0, max_documents: int = 250, staged: bool = False, data: str = "") -> dict:
    """-> {"text": what the tool says, "cut": whether max_hits cut the answer}."""
    needle = query.strip().lower()
    top = Path(root).resolve()
    base = resolve(root, subdir, data) if subdir else top
    if not base.is_dir():
        return {"text": f"not a directory: {subdir}"}
    hits: list[str] = []
    started, opened, seen, skipped = time.time(), 0, 0, 0
    for p in sorted(base.rglob("*")):
        if not p.is_file() or not _shown(top, p, staged, data) \
                or (path_contains and path_contains.lower() not in str(p).lower()):
            continue
        seen += 1
        costly = p.suffix.lower() in _DOCUMENTS
        if costly and (opened >= max_documents or time.time() - started > max_seconds):
            skipped += 1
            continue
        opened += costly
        text = extract(p)
        if text is None:
            continue
        rel = str(p.relative_to(top))
        for i, line in enumerate(text.splitlines(), 1):
            if needle in line.lower():
                said = line.strip()
                # A CSV row or a minified file is one long line, and its first 200 characters
                # read exactly like all of it. The bridged grep does not cut at all, so
                # unmarked this tool answers differently depending on whether a container is up.
                hits.append(f"{rel}:{i}: {said[:200]}"
                            + (f"  ...[+{len(said) - 200} more on this line]" if len(said) > 200 else ""))
                if len(hits) > max_hits:
                    # Past the cap the agent needs a narrower query, not an arbitrary prefix.
                    return {"text": "\n".join(hits[:max_hits]) + f"\n... more than {max_hits} matches; "
                                    "narrow the query.", "cut": True}
    # Said, not hidden: a truncated scan that reads as "no matches" sends the agent away
    # from the file it was looking for.
    note = (f"\n[scanned {seen - skipped} of {seen} files in {time.time() - started:.0f}s; "
            f"{skipped} documents were left unopened -- narrow with subdir= or path_contains=]"
            if skipped else "")
    return {"text": ("\n".join(hits) + note) if hits else f"(no line contains {query!r})" + literal(query) + note}


def near(body: str, old: str) -> str:
    """Where the passage nearly is: 12 identical failed edits in one run, each told only that the
    text was absent."""
    head = next((l.strip() for l in old.splitlines() if l.strip()), "")
    at = [n for n, l in enumerate(body.splitlines(), 1) if head and head in l][:3]
    return (f" Its first line is at line {', '.join(map(str, at))}: read there and copy from the file."
            if at else "")


def edit(root: str, path: str, old: str, new: str, data: str = "") -> str:
    p = resolve(root, path, data, write=True)
    if not p.is_file():
        return f"{path} is not a file in the workspace; list_directory shows what is."
    if p.suffix.lower() in _DOCUMENTS:
        return f"edit_file changes text files; to change {path}, send the whole new content with write_file."
    try:
        body = p.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        return f"could not read {path}: {getattr(exc, 'strerror', None) or exc}."
    hits = body.count(old)
    if hits == 0:
        return (f"that passage does not appear in {path}; read it first and copy the text "
                f"exactly, whitespace included.{near(body, old)}")
    if hits > 1:
        # Editing the first of several is how a file quietly gets the wrong one changed.
        return f"that passage appears {hits} times in {path}; extend `old` until it is unique."
    # As write does: a staged world is symlinks into the shared dataset, and writing
    # through one edits the dataset for every run after this.
    if p.is_symlink():
        p.unlink()
    p.write_text(body.replace(old, new), encoding="utf-8")
    return f"edited {p.relative_to(Path(root).resolve())} ({len(old)} chars -> {len(new)})"


OPS = {"list": listing, "read": read, "image": image, "write": write, "grep": grep, "edit": edit}


def main(call: str) -> str:
    """One operation from its JSON, as JSON: its result, or the refusal its tool raises to the model --
    a path outside, or a file the shell moved away mid-operation."""
    args = json.loads(call)
    try:
        return json.dumps({"result": OPS[args.pop("op")](**args)})
    except (ValueError, OSError) as exc:
        return json.dumps({"error": str(exc)})
